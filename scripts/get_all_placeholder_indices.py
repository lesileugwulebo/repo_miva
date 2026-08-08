import docx
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document("../Ugwulebo_Lesile_Ngozi_MIT_AWS_GCP_Updated_Network(3).docx")
pattern = re.compile(r"\[[A-Za-z0-9_/ -]{3,}\]")

print("Scanning original document for placeholders...")
for i, para in enumerate(doc.paragraphs):
    matches = pattern.findall(para.text)
    if matches:
        if not any(k in para.text for k in ["[Unit]", "[Service]", "[Install]"]):
            safe_text = para.text[:120].encode('ascii', errors='replace').decode('ascii')
            print(f"P{i}: matches={matches} text={repr(safe_text)}")
