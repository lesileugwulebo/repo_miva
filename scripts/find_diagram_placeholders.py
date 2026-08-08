import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document("../Ugwulebo_Lesile_Ngozi_MIT_AWS_GCP_Final_Thesis.docx")

print("Searching for Figure captions and placeholders...")
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if "figure" in text.lower() or "placeholder" in text.lower() or "diagram" in text.lower():
        if len(text) < 150:
            print(f"P{i}: {repr(text)}")
