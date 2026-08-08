import docx
import re

doc = docx.Document("../Ugwulebo_Lesile_Ngozi_MIT_AWS_GCP_Updated_Network(3).docx")

print("Listing placeholders with context...")
pattern = re.compile(r"\[[A-Za-z0-9_/ -]{3,}\]")

for i, para in enumerate(doc.paragraphs):
    if any(p in para.text for p in ["[INSERT", "[MET/", "[ACHIEVED", "[LOW", "[SUITABLE"]):
        print(f"P{i}: {para.text}")
